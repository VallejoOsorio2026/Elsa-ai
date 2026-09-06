"""Generación de libros XLSX sintéticos para la suite.

Los fixtures se **construyen por código**, no se versionan como archivos.
Hay dos razones y ambas importan:

- El repositorio puede ser público y ``.gitignore`` excluye ``*.xlsx``. Un
  fixture binario no podría versionarse aunque quisiéramos.
- Un archivo generado no puede contener por accidente un dato técnico real
  de planta. Todo lo que aparece aquí es inventado: códigos, subsistemas y
  descripciones no corresponden a ningún equipo de PAPELSA.

Las imágenes embebidas se inyectan escribiendo las partes OpenXML a mano
(``xl/media``, ``xl/drawings`` y sus relaciones) en vez de usar
``openpyxl.drawing``, que exigiría Pillow. El efecto secundario es
deseable: el test recorre exactamente la misma cadena de relaciones que el
parser tiene que seguir en un archivo real.
"""

import struct
import zipfile
import zlib
from io import BytesIO
from typing import Any

from defusedxml.ElementTree import fromstring
from openpyxl import Workbook

__all__ = [
    "BOM_HEADERS",
    "build_workbook",
    "minimal_png",
    "with_embedded_image",
]

_NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

BOM_HEADERS = [
    "No",
    "Subsistema",
    "Componente",
    "Código SAP",
    "Descripción técnica",
    "Cantidad",
    "Unidad",
    "Modelo/Referencia",
    "Plano de ensamble",
    "Referencia de plano",
    "Actualizar BOM",
    "Estrategia de inventario",
    "Max",
    "Min",
    "Stock Actual",
    "Observaciones",
]


def minimal_png(width: int = 4, height: int = 2) -> bytes:
    """PNG válido y diminuto, construido byte a byte (sin Pillow)."""

    def chunk(tag: bytes, payload: bytes) -> bytes:
        body = tag + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    scanlines = b"".join(b"\x00" + b"\x10\x20\x30" * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(scanlines))
        + chunk(b"IEND", b"")
    )


def build_workbook(sheets: dict[str, list[list[Any]]]) -> bytes:
    """Construye un libro con las hojas y filas indicadas."""
    workbook = Workbook()
    # `Workbook()` trae una hoja por defecto que sobra en cuanto se crean las
    # propias.
    default_sheet = workbook.active
    for index, (title, rows) in enumerate(sheets.items()):
        worksheet = default_sheet if index == 0 else workbook.create_sheet()
        worksheet.title = title
        for row in rows:
            worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_DRAWING_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    "<xdr:twoCellAnchor>"
    "<xdr:from><xdr:col>{col}</xdr:col><xdr:colOff>0</xdr:colOff>"
    "<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
    "<xdr:to><xdr:col>{col_to}</xdr:col><xdr:colOff>0</xdr:colOff>"
    "<xdr:row>{row_to}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
    "<xdr:pic><xdr:nvPicPr>"
    '<xdr:cNvPr id="1" name="Picture 1"/><xdr:cNvPicPr/>'
    "</xdr:nvPicPr><xdr:blipFill>"
    '<a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    ' r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch>'
    "</xdr:blipFill><xdr:spPr/></xdr:pic>"
    "<xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>"
)

_DRAWING_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1"'
    ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"'
    ' Target="../media/image1.png"/>'
    "</Relationships>"
)

_SHEET_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rIdDraw1"'
    ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"'
    ' Target="../drawings/drawing1.xml"/>'
    "</Relationships>"
)


def _sheet_part_for(data: bytes, sheet_name: str) -> str:
    """Ruta del XML de una hoja, siguiendo las relaciones del libro."""
    with zipfile.ZipFile(BytesIO(data)) as archive:
        workbook_root = fromstring(archive.read("xl/workbook.xml"))
        rels_root = fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        element.get("Id"): element.get("Target")
        for element in rels_root
        if element.get("Id") and element.get("Target")
    }
    for element in workbook_root.iter(f"{_NS_MAIN}sheet"):
        if element.get("name") == sheet_name:
            target = targets.get(element.get(_R_ID) or "")
            if target is None:
                break
            # Igual que en el paquete real, el destino puede ser relativo a
            # `xl/` o absoluto desde la raíz.
            if target.startswith("/"):
                return target.lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    raise LookupError(f"the workbook has no sheet named {sheet_name!r}")


def with_embedded_image(
    data: bytes,
    *,
    sheet_name: str,
    image: bytes | None = None,
    anchor_cell: tuple[int, int] = (2, 4),
) -> bytes:
    """Devuelve el libro con una imagen embebida anclada en una hoja.

    ``anchor_cell`` son índices base cero ``(columna, fila)``; el valor por
    defecto ancla en ``C5``.
    """
    image = image if image is not None else minimal_png()
    sheet_part = _sheet_part_for(data, sheet_name)
    column, row = anchor_cell

    output = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as source:
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    content = content.replace(
                        b"</Types>",
                        b'<Default Extension="png" ContentType="image/png"/>'
                        b'<Override PartName="/xl/drawings/drawing1.xml"'
                        b' ContentType="application/vnd.openxmlformats-officedocument'
                        b'.drawing+xml"/></Types>',
                    )
                elif info.filename == sheet_part:
                    # El elemento `drawing` va al final de la hoja, justo
                    # antes del cierre: el esquema exige ese orden.
                    content = content.replace(
                        b"</worksheet>",
                        b'<drawing xmlns:r="http://schemas.openxmlformats.org/'
                        b'officeDocument/2006/relationships" r:id="rIdDraw1"/></worksheet>',
                    )
                target.writestr(info, content)

            directory, filename = sheet_part.rsplit("/", 1)
            target.writestr(f"{directory}/_rels/{filename}.rels", _SHEET_RELS)
            target.writestr(
                "xl/drawings/drawing1.xml",
                _DRAWING_XML.format(col=column, row=row, col_to=column + 4, row_to=row + 10),
            )
            target.writestr("xl/drawings/_rels/drawing1.xml.rels", _DRAWING_RELS)
            target.writestr("xl/media/image1.png", image)
    return output.getvalue()
