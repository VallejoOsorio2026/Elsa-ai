"""Lectura del XLSX aprobado por Ingeniería.

El archivo representa *cómo debía quedar alimentado el BOM* y cuál es la
información técnica de referencia. Se conserva como fuente histórica y se
interpreta aquí, sin ejecutar nada de lo que traiga dentro.

Qué NO hace este parser, y por qué:

- **No evalúa fórmulas.** El libro se abre con ``data_only=False``, de modo
  que una celda con fórmula llega como texto (``"=A1*2"``) y se trata como
  valor no interpretable. Leer el resultado que Excel dejó cacheado sería
  aceptar como verdad técnica un número que nadie en ELSA puede verificar.
  El NPR del AMEF se recalcula siempre en el backend.
- **No sigue enlaces externos.** El libro se abre con ``keep_links=False``:
  los vínculos se descartan al cargar y no se resuelve ninguna dirección.
- **No interpreta los planos.** Las imágenes se extraen como evidencia. No
  hay OCR ni visión artificial en este bloque.

Cómo tolera un archivo que no es exactamente el previsto: las hojas y las
columnas se reconocen por **sinónimos normalizados**, no por posición ni por
nombre exacto. Una columna que no se reconoce no se pierde: va a ``extra``
con un aviso. Una asociación que no es segura no se inventa: queda
pendiente de revisión.
"""

import logging
import re
from collections.abc import Sequence
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from elsa.core.normalization import (
    normalize_key,
    normalize_quantity,
    normalize_sap_code,
    normalize_text,
)
from elsa.ingestion.errors import MissingSheetError, UnrecognizedFormatError
from elsa.ingestion.model import (
    IngestionWarning,
    ParsedBomRow,
    ParsedDrawingImage,
    ParsedEngineeringBom,
    ParsedFailureMode,
    ParsedOptionRow,
    ParsedSodCriterion,
)
from elsa.ingestion.safety import ArchiveLimits, inspect_xlsx
from elsa.ingestion.xlsx_drawings import extract_embedded_images
from elsa.ports.artifact_storage import sha256_hex

_logger = logging.getLogger("elsa.ingestion.xlsx")

__all__ = ["parse_engineering_workbook"]

# Cuántas filas se miran buscando la fila de encabezados. Las plantillas
# suelen traer título, logotipo y filas en blanco antes de la tabla.
_HEADER_SCAN_ROWS = 30

# Encabezados reconocidos que debe tener una fila para considerarse la fila
# de encabezados. Con menos, cualquier fila de datos podría confundirse.
_MIN_HEADER_MATCHES = 3

# Tope defensivo de filas. El límite de tamaño del paquete ya acota el
# archivo; esto evita que una hoja con un rango declarado absurdo obligue a
# recorrer millones de filas vacías.
_MAX_ROWS = 200_000

# --- Hojas -----------------------------------------------------------

_SHEET_SYNONYMS: dict[str, tuple[str, ...]] = {
    "bom": ("bom", "lista de materiales", "listado bom", "despiece bom"),
    "drawings": ("plano despiece", "planos", "despiece", "plano de despiece"),
    "options": ("tablas de opciones (bom)", "tablas de opciones", "opciones"),
    "amef": ("amef", "fmea", "analisis de modos de falla", "modos de falla"),
    "sod": ("valoracion sod", "valoracion s/o/d", "sod", "criterios sod", "valoracion"),
}

# --- Columnas --------------------------------------------------------


def _synonym_index(definitions: dict[str, tuple[str, ...]]) -> dict[str, str]:
    """Invierte campo → sinónimos en sinónimo normalizado → campo."""
    index: dict[str, str] = {}
    for field, synonyms in definitions.items():
        for synonym in synonyms:
            key = normalize_key(synonym)
            if key is not None:
                index.setdefault(key, field)
    return index


_BOM_COLUMNS: dict[str, tuple[str, ...]] = {
    "position": ("no", "n", "num", "numero", "item", "posicion", "pos", "#", "no."),
    "subsystem_name": ("subsistema", "sistema", "conjunto", "subconjunto"),
    "component_name": ("componente", "elemento", "pieza", "nombre del componente"),
    "sap_code": ("codigo sap", "cod sap", "codigo", "material", "material sap", "sap"),
    "technical_description": (
        "descripcion tecnica",
        "descripcion",
        "especificacion",
        "detalle tecnico",
    ),
    "quantity": ("cantidad", "cant", "cant.", "qty", "ctd"),
    "unit": ("unidad", "um", "u.m.", "udm", "unidad de medida"),
    "model_reference": (
        "modelo/referencia",
        "modelo referencia",
        "modelo",
        "referencia",
        "part number",
        "pn",
        "numero de parte",
    ),
    "assembly_drawing": (
        "plano de ensamble",
        "plano ensamble",
        "plano",
        "numero de plano",
        "no plano",
    ),
    "drawing_reference": (
        "referencia de plano",
        "referencia plano",
        "ref plano",
        "item plano",
        "referencia",
    ),
    "bom_update_flag": ("actualizar bom", "actualizacion bom", "actualiza bom", "bom actualizado"),
    "inventory_strategy": (
        "estrategia de inventario",
        "estrategia inventario",
        "estrategia",
    ),
    "stock_max": ("max", "maximo", "stock max", "stock maximo"),
    "stock_min": ("min", "minimo", "stock min", "stock minimo"),
    "source_stock": ("stock actual", "stock", "existencia", "existencias"),
    "remarks": ("observaciones", "observacion", "notas", "comentarios"),
}

_AMEF_COLUMNS: dict[str, tuple[str, ...]] = {
    "subsystem_name": ("subsistema", "sistema", "conjunto"),
    "component_name": ("componente", "elemento", "pieza"),
    "sap_code": ("codigo sap", "cod sap", "codigo", "material"),
    "failure_mode": ("modo de falla", "modo falla", "modo de fallo", "modo"),
    "effect": ("efecto", "efecto de falla", "efecto de la falla"),
    "cause": ("causa", "causa de falla", "causa de la falla"),
    "severity": ("s", "severidad", "sev"),
    "occurrence": ("o", "ocurrencia", "ocu", "frecuencia"),
    "detection": ("d", "deteccion", "det"),
    "reported_rpn": ("npr", "rpn", "nrp"),
    "action": ("accion a tomar", "accion", "acciones", "accion recomendada"),
    "preventive_plan": ("plan preventivo", "preventivo", "plan de mantenimiento preventivo"),
    "corrective_action": ("accion correctiva", "correctivo", "correctiva"),
    "remarks": ("observaciones", "observacion", "notas", "comentarios"),
}

_BOM_INDEX = _synonym_index(_BOM_COLUMNS)
_AMEF_INDEX = _synonym_index(_AMEF_COLUMNS)

# Vocabulario de las tres dimensiones. Se compara por subcadena sobre el
# texto normalizado, así que cubre «Severidad», «Índice de severidad» o
# «Valoración de la severidad» sin enumerarlos.
_SOD_DIMENSIONS: dict[str, str] = {
    "severidad": "severity",
    "gravedad": "severity",
    "severity": "severity",
    "ocurrencia": "occurrence",
    "occurrence": "occurrence",
    "frecuencia": "occurrence",
    "probabilidad": "occurrence",
    "deteccion": "detection",
    "detection": "detection",
}

# Rótulos que designan la dimensión por su inicial, como «Calificación (S)»
# o simplemente «(O)». Se buscan como token aislado para que la «s» de
# cualquier palabra no dispare una coincidencia.
_SOD_INITIALS: dict[str, str] = {"s": "severity", "o": "occurrence", "d": "detection"}
_SOD_INITIAL_PATTERN = re.compile(r"\(\s*([sod])\s*\)")

# Rótulos de la columna que lleva el valor de la escala. Cuando la tabla
# declara uno, manda él: elegir «el primer número de la fila» falla en cuanto
# la tabla antepone una tasa o un porcentaje, que es justo lo que suelen
# hacer las tablas de ocurrencia.
_SOD_SCALE_COLUMNS: tuple[str, ...] = (
    "valor",
    "escala",
    "nivel",
    "indice",
    "calificacion",
    "puntaje",
    "rango",
    "grado",
    "criterio numerico",
)


def _dimension_of(text: str | None) -> str | None:
    """Dimensión S/O/D que designa un rótulo, o ``None``."""
    key = normalize_key(text)
    if key is None:
        return None
    match = _SOD_INITIAL_PATTERN.search(key)
    if match:
        return _SOD_INITIALS[match.group(1)]
    for name, dimension in _SOD_DIMENSIONS.items():
        if name in key:
            return dimension
    # Un rótulo que es exactamente la inicial («S», «O», «D») también vale.
    return _SOD_INITIALS.get(key.strip())


# Longitud mínima de un sinónimo para admitir coincidencia laxa. Debajo de
# esto (`s`, `o`, `d`, `no`) una coincidencia parcial sería ruido.
_MIN_LOOSE_SYNONYM = 4


@dataclass(frozen=True, slots=True)
class _ColumnMap:
    """Correspondencia entre columnas de la hoja y campos del modelo."""

    header_row: int
    by_field: dict[str, int]
    extras: dict[int, str]
    warnings: tuple[IngestionWarning, ...]

    def value(self, row: Sequence[Any], field: str) -> Any:
        index = self.by_field.get(field)
        if index is None or index >= len(row):
            return None
        return row[index]


def _map_columns(header: Sequence[Any], index: dict[str, str], sheet: str) -> _ColumnMap:
    """Empareja encabezados con campos, primero exacto y luego laxo."""
    by_field: dict[str, int] = {}
    extras: dict[int, str] = {}
    warnings: list[IngestionWarning] = []
    loose: list[str] = []
    unknown: list[str] = []

    for position, cell in enumerate(header):
        original = normalize_text(cell)
        if original is None:
            continue
        key = normalize_key(original)
        if key is None:
            continue

        field = index.get(key)
        if field is None:
            # Coincidencia laxa: el encabezado real suele traer unidades o
            # aclaraciones («Cantidad (UN)», «Código SAP - material»). Solo se
            # acepta si apunta a un único campo; si apunta a varios, es
            # ambigua y el valor se conserva sin interpretar.
            candidates = {
                candidate_field
                for synonym, candidate_field in index.items()
                if len(synonym) >= _MIN_LOOSE_SYNONYM and (synonym in key or key in synonym)
            }
            if len(candidates) == 1:
                field = candidates.pop()
                loose.append(original)

        if field is None or field in by_field:
            # Una columna repetida no se sobrescribe: la primera manda y la
            # segunda se conserva como extra, para no perder su contenido.
            extras[position] = original
            if field is None:
                unknown.append(original)
            continue
        by_field[field] = position

    if loose:
        warnings.append(
            IngestionWarning(
                code="column_matched_loosely",
                message="Some column headers matched a known field only approximately.",
                location=f"{sheet}: {', '.join(sorted(loose)[:8])}",
            )
        )
    if unknown:
        warnings.append(
            IngestionWarning(
                code="unknown_columns",
                message="Some columns were not recognised and are preserved as extra data.",
                location=f"{sheet}: {', '.join(sorted(unknown)[:8])}",
            )
        )
    return _ColumnMap(header_row=0, by_field=by_field, extras=extras, warnings=tuple(warnings))


def _find_header_row(rows: Sequence[Sequence[Any]], index: dict[str, str]) -> int | None:
    """Índice de la fila con más encabezados reconocidos exactamente."""
    best_row: int | None = None
    best_score = 0
    for position, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        score = 0
        seen: set[str] = set()
        for cell in row:
            key = normalize_key(cell)
            field = index.get(key) if key else None
            if field is not None and field not in seen:
                seen.add(field)
                score += 1
        if score > best_score:
            best_score, best_row = score, position
    return best_row if best_score >= _MIN_HEADER_MATCHES else None


# --- Utilidades de celda ---------------------------------------------


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _original_text(value: Any) -> str | None:
    """Representación textual del valor tal y como venía en la celda."""
    if value is None:
        return None
    if isinstance(value, str):
        return normalize_text(value)
    return str(value)


def _scale_value(value: Any) -> int | None:
    """Un factor S/O/D válido: entero de 1 a 10. Cualquier otra cosa, ``None``."""
    quantity = normalize_quantity(value)
    if quantity is None or quantity != quantity.to_integral_value():
        return None
    number = int(quantity)
    return number if 1 <= number <= 10 else None


def _row_is_empty(row: Sequence[Any]) -> bool:
    return all(normalize_text(cell) is None for cell in row)


def _extras(row: Sequence[Any], columns: _ColumnMap) -> dict[str, str]:
    values: dict[str, str] = {}
    for position, header in columns.extras.items():
        if position < len(row):
            text = _original_text(row[position])
            if text is not None:
                values[header] = text
    return values


# --- Hojas concretas -------------------------------------------------


def _resolve_sheets(titles: Sequence[str]) -> dict[str, str]:
    """Empareja los títulos reales del libro con los roles conocidos."""
    resolved: dict[str, str] = {}
    normalized = {title: normalize_key(title) or "" for title in titles}
    for role, synonyms in _SHEET_SYNONYMS.items():
        keys = [normalize_key(synonym) or "" for synonym in synonyms]
        for title, key in normalized.items():
            if title in resolved.values():
                continue
            if key in keys or any(candidate and candidate in key for candidate in keys):
                resolved[role] = title
                break
    return resolved


def _parse_bom(
    rows: Sequence[Sequence[Any]], sheet: str
) -> tuple[tuple[ParsedBomRow, ...], tuple[IngestionWarning, ...]]:
    header_row = _find_header_row(rows, _BOM_INDEX)
    if header_row is None:
        raise UnrecognizedFormatError(f"the '{sheet}' sheet has no recognisable header row")
    columns = _map_columns(rows[header_row], _BOM_INDEX, sheet)
    warnings = list(columns.warnings)
    formulas = 0
    parsed: list[ParsedBomRow] = []

    for offset, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        if _row_is_empty(row):
            continue
        raw_quantity = columns.value(row, "quantity")
        if _is_formula(raw_quantity):
            formulas += 1
        sap_raw = columns.value(row, "sap_code")
        parsed.append(
            ParsedBomRow(
                source_row=offset,
                position=_original_text(columns.value(row, "position")),
                subsystem_name=normalize_text(columns.value(row, "subsystem_name")),
                component_name=normalize_text(columns.value(row, "component_name")),
                sap_code=normalize_sap_code(sap_raw),
                sap_code_original=_original_text(sap_raw),
                technical_description=normalize_text(columns.value(row, "technical_description")),
                quantity=normalize_quantity(raw_quantity),
                quantity_original=_original_text(raw_quantity),
                unit=normalize_text(columns.value(row, "unit")),
                model_reference=normalize_text(columns.value(row, "model_reference")),
                assembly_drawing=normalize_text(columns.value(row, "assembly_drawing")),
                drawing_reference=normalize_text(columns.value(row, "drawing_reference")),
                bom_update_flag=normalize_text(columns.value(row, "bom_update_flag")),
                inventory_strategy=normalize_text(columns.value(row, "inventory_strategy")),
                stock_max=normalize_quantity(columns.value(row, "stock_max")),
                stock_min=normalize_quantity(columns.value(row, "stock_min")),
                source_stock=normalize_quantity(columns.value(row, "source_stock")),
                remarks=normalize_text(columns.value(row, "remarks")),
                extra=_extras(row, columns),
            )
        )

    if formulas:
        warnings.append(
            IngestionWarning(
                code="formula_not_evaluated",
                message=(
                    "Some quantities are formulas. They are preserved as written and left "
                    "without a numeric value: the workbook is not evaluated."
                ),
                location=f"{sheet}: {formulas} row(s)",
            )
        )
    return tuple(parsed), tuple(warnings)


def _parse_amef(
    rows: Sequence[Sequence[Any]], sheet: str
) -> tuple[tuple[ParsedFailureMode, ...], tuple[IngestionWarning, ...]]:
    header_row = _find_header_row(rows, _AMEF_INDEX)
    if header_row is None:
        return (), (
            IngestionWarning(
                code="unrecognised_sheet",
                message="The AMEF sheet has no recognisable header row and was not imported.",
                location=sheet,
            ),
        )
    columns = _map_columns(rows[header_row], _AMEF_INDEX, sheet)
    warnings = list(columns.warnings)
    disagreements = 0
    parsed: list[ParsedFailureMode] = []

    for offset, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        if _row_is_empty(row):
            continue
        failure_mode = ParsedFailureMode(
            source_row=offset,
            subsystem_name=normalize_text(columns.value(row, "subsystem_name")),
            component_name=normalize_text(columns.value(row, "component_name")),
            sap_code=normalize_sap_code(columns.value(row, "sap_code")),
            failure_mode=normalize_text(columns.value(row, "failure_mode")),
            effect=normalize_text(columns.value(row, "effect")),
            cause=normalize_text(columns.value(row, "cause")),
            severity=_scale_value(columns.value(row, "severity")),
            occurrence=_scale_value(columns.value(row, "occurrence")),
            detection=_scale_value(columns.value(row, "detection")),
            action=normalize_text(columns.value(row, "action")),
            preventive_plan=normalize_text(columns.value(row, "preventive_plan")),
            corrective_action=normalize_text(columns.value(row, "corrective_action")),
            remarks=normalize_text(columns.value(row, "remarks")),
            extra=_extras(row, columns),
        )
        # El NPR que traiga el archivo se compara, nunca se usa. Si difiere del
        # calculado, manda el calculado y queda constancia de la diferencia.
        reported = normalize_quantity(columns.value(row, "reported_rpn"))
        if (
            reported is not None
            and failure_mode.rpn is not None
            and reported != Decimal(failure_mode.rpn)
        ):
            disagreements += 1
        parsed.append(failure_mode)

    if disagreements:
        warnings.append(
            IngestionWarning(
                code="rpn_recomputed",
                message=(
                    "The RPN written in the workbook differs from S x O x D. The value "
                    "computed by the backend prevails."
                ),
                location=f"{sheet}: {disagreements} row(s)",
            )
        )
    return tuple(parsed), tuple(warnings)


def _sod_announcements(row: Sequence[Any]) -> list[tuple[int, str]]:
    """Dimensiones que anuncia una fila, con la columna en la que lo hace."""
    found: list[tuple[int, str]] = []
    for position, cell in enumerate(row):
        dimension = _dimension_of(normalize_text(cell))
        if dimension is not None and dimension not in {name for _, name in found}:
            found.append((position, dimension))
    return found


def _sod_scale_columns(row: Sequence[Any]) -> dict[int, int]:
    """Columnas de escala declaradas por una fila de rótulos.

    Devuelve, por cada columna de escala encontrada, su propia posición. Se
    usa para leer el valor de la escala de la columna correcta en vez de
    tomar «el primer número de la fila».
    """
    columns: dict[int, int] = {}
    for position, cell in enumerate(row):
        key = normalize_key(cell)
        if key is None:
            continue
        if any(name in key for name in _SOD_SCALE_COLUMNS):
            columns[position] = position
    return columns


def _dimension_for_column(
    position: int, layout: list[tuple[int, str]], stacked: str | None
) -> str | None:
    """Dimensión a la que pertenece una columna.

    Con las tres tablas una debajo de otra, la dimensión es la última
    anunciada y no depende de la columna. Con las tablas una al lado de otra,
    cada una gobierna desde la columna en la que se anuncia hasta donde
    empieza la siguiente.
    """
    if len(layout) > 1:
        current: str | None = None
        for start, dimension in layout:
            if position >= start:
                current = dimension
        return current
    return stacked


def _parse_sod(
    rows: Sequence[Sequence[Any]], sheet: str
) -> tuple[tuple[ParsedSodCriterion, ...], tuple[IngestionWarning, ...]]:
    """Lee los criterios S/O/D.

    Se admiten las dos formas en que la plantilla puede disponer las tres
    tablas, porque ambas son igual de habituales y elegir solo una hacía que
    la otra se perdiera **en silencio**:

    - **Apiladas**: cada tabla va precedida de su rótulo y se recorre de
      arriba abajo recordando la última dimensión anunciada.
    - **Una al lado de otra**: una misma fila anuncia varias dimensiones en
      columnas distintas, y cada una gobierna su franja de columnas.

    El valor de la escala se toma de la columna que la tabla declare
    («Valor», «Nivel», «Calificación»…). Solo si no declara ninguna se
    recurre al primer entero entre 1 y 10 de la fila.

    Una dimensión anunciada que no produce ningún criterio genera un aviso:
    quedarse callado ahí es exactamente el defecto que se corrigió.
    """
    criteria: list[ParsedSodCriterion] = []
    warnings: list[IngestionWarning] = []

    stacked: str | None = None
    layout: list[tuple[int, str]] = []
    scale_columns: dict[int, int] = {}
    announced: set[str] = set()
    orphan_rows = 0

    for offset, row in enumerate(rows, start=1):
        if _row_is_empty(row):
            continue

        announcements = _sod_announcements(row)
        if announcements:
            if len(announcements) > 1:
                # Varias dimensiones en la misma fila: tablas en paralelo.
                layout = announcements
            else:
                stacked = announcements[0][1]
            announced.update(dimension for _, dimension in announcements)
            # La fila que anuncia puede además declarar columnas de escala.
            scale_columns = _sod_scale_columns(row) or scale_columns
            continue

        declared = _sod_scale_columns(row)
        if declared:
            scale_columns = declared
            continue

        # Candidatos a valor de escala: los de las columnas declaradas, o
        # cualquier entero de 1 a 10 si la tabla no declaró ninguna.
        positions = sorted(scale_columns) if scale_columns else range(len(row))
        for position in positions:
            if position >= len(row):
                continue
            scale = _scale_value(row[position])
            if scale is None:
                continue
            dimension = _dimension_for_column(position, layout, stacked)
            if dimension is None:
                orphan_rows += 1
                continue

            # El texto y los rangos se leen de la franja de columnas de esta
            # dimensión, para que dos tablas en paralelo no se mezclen.
            limit = len(row)
            for start, _ in layout:
                if start > position:
                    limit = min(limit, start)
            window = list(row[position:limit])

            remaining = [
                normalize_text(cell)
                for index, cell in enumerate(window)
                if index != 0 and normalize_text(cell) is not None
            ]
            numbers = [
                value
                for index, cell in enumerate(window)
                if index != 0 and (value := normalize_quantity(cell)) is not None
            ]
            criteria.append(
                ParsedSodCriterion(
                    source_row=offset,
                    dimension=dimension,
                    scale_value=scale,
                    label=remaining[0] if remaining else None,
                    description=remaining[1] if len(remaining) > 1 else None,
                    range_low=numbers[0] if numbers else None,
                    range_high=numbers[1] if len(numbers) > 1 else None,
                )
            )
            if not scale_columns:
                # Sin columna declarada, el primer entero válido de la fila
                # es el valor: seguir buscando duplicaría el criterio.
                break

    if orphan_rows:
        warnings.append(
            IngestionWarning(
                code="sod_dimension_unknown",
                message=(
                    "Some S/O/D rows appear before any dimension heading and were not "
                    "imported: the dimension is not guessed."
                ),
                location=f"{sheet}: {orphan_rows} row(s)",
            )
        )

    extracted = {criterion.dimension for criterion in criteria}
    empty = sorted(announced - extracted)
    if empty:
        warnings.append(
            IngestionWarning(
                code="sod_dimension_without_criteria",
                message=(
                    "A S/O/D dimension is announced in the sheet but no scale value could "
                    "be read for it. The criteria are not invented."
                ),
                location=f"{sheet}: {', '.join(empty)}",
            )
        )
    missing = sorted({"severity", "occurrence", "detection"} - extracted)
    if missing:
        warnings.append(
            IngestionWarning(
                code="sod_dimension_missing",
                message="The S/O/D sheet did not yield criteria for every dimension.",
                location=f"{sheet}: {', '.join(missing)}",
            )
        )
    return tuple(criteria), tuple(warnings)


def _parse_options(rows: Sequence[Sequence[Any]], sheet: str) -> tuple[ParsedOptionRow, ...]:
    """Lee las tablas de opciones: una columna por tabla, encabezada por su nombre."""
    if not rows:
        return ()
    header = rows[0]
    names = {
        position: text
        for position, cell in enumerate(header)
        if (text := normalize_text(cell)) is not None
    }
    if not names:
        return ()
    options: list[ParsedOptionRow] = []
    for offset, row in enumerate(rows[1:], start=2):
        for position, table_name in names.items():
            if position >= len(row):
                continue
            value = normalize_text(row[position])
            if value is None:
                continue
            options.append(
                ParsedOptionRow(source_row=offset, table_name=table_name, option_value=value)
            )
    return tuple(options)


def _associate_drawings(
    images: Sequence[Any],
    sheet_values: dict[str, set[str]],
    declared_drawings: set[str],
) -> tuple[tuple[ParsedDrawingImage, ...], tuple[IngestionWarning, ...]]:
    """Asocia cada imagen a un número de plano, solo cuando es inequívoco.

    Regla única y conservadora: si la hoja donde está anclada la imagen
    menciona **exactamente un** número de plano de los declarados en el BOM,
    la asociación es esa. Con cero o con varios, no se elige: la imagen se
    conserva como evidencia y queda pendiente de revisión.

    Cualquier regla más ambiciosa exigiría mirar el contenido gráfico, que es
    precisamente lo que este bloque no hace.
    """
    parsed: list[ParsedDrawingImage] = []
    pending = 0

    for image in images:
        candidates: set[str] = set()
        if image.sheet_name is not None:
            candidates = sheet_values.get(image.sheet_name, set()) & declared_drawings

        number = candidates.pop() if len(candidates) == 1 else None
        if number is None:
            pending += 1
        parsed.append(
            ParsedDrawingImage(
                content=image.content,
                content_type=image.content_type,
                sha256=sha256_hex(image.content),
                sheet_name=image.sheet_name,
                anchor=image.anchor,
                width_px=image.width_px,
                height_px=image.height_px,
                drawing_number=number,
                association_rule="unique_drawing_number_in_sheet" if number else None,
            )
        )

    warnings: list[IngestionWarning] = []
    if pending:
        warnings.append(
            IngestionWarning(
                code="drawing_association_pending",
                message=(
                    "Some embedded images could not be tied to a single drawing number. "
                    "They are kept as evidence and left for human review; the association "
                    "is never guessed."
                ),
                location=f"{pending} image(s)",
            )
        )
    return tuple(parsed), tuple(warnings)


# --- Entrada pública -------------------------------------------------


def parse_engineering_workbook(
    data: bytes, *, limits: ArchiveLimits | None = None
) -> ParsedEngineeringBom:
    """Interpreta el XLSX de Ingeniería y devuelve sus datos.

    Revisa primero la seguridad del paquete (:func:`inspect_xlsx`) y solo
    entonces lee contenido. Lanza una subclase de
    :class:`~elsa.ingestion.errors.IngestionError` si el archivo no puede
    aceptarse o su hoja de BOM no puede interpretarse.
    """
    inspection = inspect_xlsx(data, limits)
    warnings: list[IngestionWarning] = list(inspection.warnings)

    # `keep_links=False` descarta los vínculos externos al cargar; `read_only`
    # evita construir el modelo completo del libro en memoria.
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        titles = tuple(workbook.sheetnames)
        roles = _resolve_sheets(titles)
        if "bom" not in roles:
            raise MissingSheetError(
                "the workbook has no BOM sheet; the engineering source cannot be interpreted"
            )

        sheet_rows: dict[str, list[list[Any]]] = {}
        sheet_values: dict[str, set[str]] = {}
        for title in titles:
            rows: list[list[Any]] = []
            values: set[str] = set()
            for row in workbook[title].iter_rows(values_only=True):
                if len(rows) >= _MAX_ROWS:
                    warnings.append(
                        IngestionWarning(
                            code="sheet_truncated",
                            message="The sheet exceeds the maximum number of rows read.",
                            location=title,
                        )
                    )
                    break
                cells = list(row)
                rows.append(cells)
                for cell in cells:
                    text = normalize_text(cell)
                    if text is not None:
                        values.add(text)
            sheet_rows[title] = rows
            sheet_values[title] = values

        images = extract_embedded_images(data)
    finally:
        workbook.close()

    bom_rows, bom_warnings = _parse_bom(sheet_rows[roles["bom"]], roles["bom"])
    warnings.extend(bom_warnings)

    failure_modes: tuple[ParsedFailureMode, ...] = ()
    if "amef" in roles:
        failure_modes, amef_warnings = _parse_amef(sheet_rows[roles["amef"]], roles["amef"])
        warnings.extend(amef_warnings)
    else:
        warnings.append(
            IngestionWarning(
                code="missing_sheet",
                message="The workbook has no AMEF sheet; no failure modes were imported.",
                location="amef",
            )
        )

    sod_criteria: tuple[ParsedSodCriterion, ...] = ()
    if "sod" in roles:
        sod_criteria, sod_warnings = _parse_sod(sheet_rows[roles["sod"]], roles["sod"])
        warnings.extend(sod_warnings)
    else:
        warnings.append(
            IngestionWarning(
                code="missing_sheet",
                message="The workbook has no S/O/D sheet; no criteria were imported.",
                location="sod",
            )
        )

    option_rows: tuple[ParsedOptionRow, ...] = ()
    if "options" in roles:
        option_rows = _parse_options(sheet_rows[roles["options"]], roles["options"])

    declared_drawings = {
        row.assembly_drawing for row in bom_rows if row.assembly_drawing is not None
    }
    drawing_images, drawing_warnings = _associate_drawings(images, sheet_values, declared_drawings)
    warnings.extend(drawing_warnings)

    _logger.info(
        "engineering workbook parsed",
        extra={
            "bom_rows": len(bom_rows),
            "failure_modes": len(failure_modes),
            "sod_criteria": len(sod_criteria),
            "drawings": len(drawing_images),
            "warnings": len(warnings),
        },
    )
    return ParsedEngineeringBom(
        sheets_detected=titles,
        bom_rows=bom_rows,
        failure_modes=failure_modes,
        sod_criteria=sod_criteria,
        option_rows=option_rows,
        drawing_images=drawing_images,
        warnings=tuple(warnings),
    )
